"""گزارش روزانهٔ اکسل و ارسال با ایمیل (Gmail).

روی سرور خودکار است (bot/web.py):
- سرویس رایگان cron-job.org هر شب ساعت ۲۳:۵۵ تهران آدرس /report?key=... را باز می‌کند.
- اگر آن اجرا نشد، اولین پیام بعد از نیمه‌شب گزارش دیروز را می‌فرستد.

اجرای دستی:
    cd ~/my-pro && python3.11 -m bot.report

گزینه‌ها:
    --date 2026-09-30   گزارش یک روز دیگر (میلادی)
    --dry-run           فقط فایل اکسل را بساز، ایمیل نفرست

اطلاعات ایمیل در فایل mail.txt کنار پروژه (دو خط):
    آدرس جیمیل فرستنده
    App Password شانزده حرفی گوگل
"""

import argparse
import io
import json
import os
import smtplib
from datetime import datetime, timedelta
from email.message import EmailMessage

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .pin import gregorian_to_jalali
from .storage import ROOT, TEHRAN, records_for_day

REPORT_TO = os.environ.get("REPORT_TO", "alifitnes137012@gmail.com")
FLOW_NAMES = {"pin": "رمز عابر بانکی", "zikr": "ذکر شخصی"}

COLUMNS = [
    ("ردیف", 7),
    ("ساعت", 9),
    ("سرویس", 16),
    ("نام و نام خانوادگی", 24),
    ("شماره همراه", 15),
    ("اسم صدا زده‌شده", 16),
    ("تاریخ تولد", 12),
    ("نام پدر", 14),
    ("نام مادر", 14),
    ("آیدی تلگرام", 16),
]


def _jalali(day: str) -> str:
    y, m, d = map(int, day.split("-"))
    jy, jm, jd = gregorian_to_jalali(y, m, d)
    return f"{jy}/{jm:02d}/{jd:02d}"


def build_excel(day: str, rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = _jalali(day).replace("/", "-")
    ws.sheet_view.rightToLeft = True

    gold = PatternFill("solid", fgColor="F0C86E")
    thin = Side(style="thin", color="C9B27A")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    ws.append([name for name, _ in COLUMNS])
    for i, (_, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill, cell.alignment, cell.border = Font(bold=True, size=12), gold, center, border
        ws.column_dimensions[cell.column_letter].width = width

    for n, r in enumerate(rows, start=1):
        date = r.get("date")
        birth = f"{date[2]}/{date[1]:02d}/{date[0]:02d}" if date else ""
        who = f"@{r['username']}" if r.get("username") else str(r.get("user_id") or "")
        ws.append([
            n,
            r["created"][11:16],
            FLOW_NAMES.get(r["flow"], r["flow"]),
            r.get("full", ""),
            r.get("phone", ""),
            r.get("called", ""),
            birth,
            r.get("father", ""),
            r.get("mother", ""),
            who,
        ])
        for cell in ws[ws.max_row]:
            cell.alignment, cell.border = center, border
        ws.cell(row=ws.max_row, column=5).number_format = "@"  # صفر اول شماره حفظ شود

    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def send_email(day: str, rows: list[dict], xlsx: bytes) -> None:
    sender, password = (ROOT / "mail.txt").read_text().split()[:2]
    jday = _jalali(day)
    msg = EmailMessage()
    msg["Subject"] = f"گزارش روزانهٔ مینی اپ — {jday} — {len(rows)} نفر"
    msg["From"] = sender
    msg["To"] = REPORT_TO
    msg.set_content(
        f"گزارش روز {jday}\nتعداد مراجعه‌ها: {len(rows)}\n\nفایل اکسل پیوست است.\n"
        "مینی اپ استاد فاطمه سادات جعفرنیا"
    )
    msg.add_attachment(
        xlsx,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"report-{jday.replace('/', '-')}.xlsx",
    )
    with smtplib.SMTP("smtp.gmail.com", 587, timeout=60) as smtp:
        smtp.starttls()
        smtp.login(sender, password)
        smtp.send_message(msg)


STATE = ROOT / "report_state.json"


def sent_days() -> set[str]:
    return set(json.loads(STATE.read_text())) if STATE.exists() else set()


def send_report(day: str) -> int:
    """گزارش یک روز را می‌فرستد و در report_state.json ثبت می‌کند که فرستاده شده."""
    rows = records_for_day(day)
    send_email(day, rows, build_excel(day, rows))
    STATE.write_text(json.dumps(sorted(sent_days() | {day})))
    return len(rows)


def yesterday() -> str:
    return (datetime.now(TEHRAN) - timedelta(days=1)).strftime("%Y-%m-%d")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--date", default=datetime.now(TEHRAN).strftime("%Y-%m-%d"))
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    rows = records_for_day(args.date)
    xlsx = build_excel(args.date, rows)
    if args.dry_run:
        path = ROOT / f"report-{args.date}.xlsx"
        path.write_bytes(xlsx)
        print(f"{len(rows)} rows -> {path}")
        return
    print(f"✅ sent {send_report(args.date)} rows to {REPORT_TO}")


if __name__ == "__main__":
    main()
